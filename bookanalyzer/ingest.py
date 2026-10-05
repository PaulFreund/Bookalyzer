"""Reproducible text import for EPUB, DOCX, TXT, and text-based PDF files."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub
from pypdf import PdfReader

from .config import project_path
from .provenance import sha256_file, sha256_text, utc_now


SUPPORTED_EXTENSIONS = {".epub", ".docx", ".txt", ".pdf"}
WORD_PATTERN = re.compile(r"[^\W_]+(?:[’'\-][^\W_]+)*", flags=re.UNICODE)
CHAPTER_PATTERN = re.compile(
    r"^(?:chapter|kapitel|part|teil|book)\s+(?:[\divxlcdm]+|[\w-]+)(?:\s*[:.\-—–]\s*.*)?$",
    flags=re.IGNORECASE,
)
SCENE_PATTERN = re.compile(r"^\s*(?:\*\s*){3,}$|^\s*(?:#\s*){3,}$|^\s*[-—]\s*[-—]\s*[-—]\s*$")
PAGE_NUMBER_PATTERN = re.compile(r"^\s*(?:page\s+)?\d+(?:\s+of\s+\d+)?\s*$", re.IGNORECASE)

FRONTMATTER_HEADINGS = {
    "contents",
    "table of contents",
    "inhaltsverzeichnis",
    "copyright",
    "imprint",
    "impressum",
    "title page",
    "dedication",
    "widmung",
    "foreword",
    "vorwort",
}
BACKMATTER_HEADINGS = {
    "acknowledgements",
    "acknowledgments",
    "danksagung",
    "about the author",
    "über den autor",
    "über die autorin",
    "bibliography",
    "bibliografie",
    "also by",
    "advertisement",
    "advertisements",
}


class IngestError(RuntimeError):
    """Raised when an input cannot be safely imported."""


@dataclass
class ExtractedSection:
    chapter_title: str
    content: str
    source_position: str


def count_words(text: str) -> int:
    return len(WORD_PATTERN.findall(text))


def normalize_text(text: str) -> str:
    """Apply only loss-minimizing, specified normalization."""
    normalized = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    normalized = re.sub(r"[ \t]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _slug(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")
    return slug or "book"


def _read_txt(path: Path) -> tuple[str, list[ExtractedSection], list[dict[str, Any]]]:
    raw = path.read_bytes()
    decoded: str | None = None
    encoding = ""
    for candidate in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            decoded = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise IngestError(f"Could not decode text file as UTF-8 or CP1252: {path}")
    sections = _split_plain_text(decoded)
    log = [{"operation": "decode", "detail": {"encoding": encoding}}]
    return path.stem, sections, log


def _split_plain_text(text: str) -> list[ExtractedSection]:
    lines = text.splitlines()
    sections: list[ExtractedSection] = []
    title = "Text"
    buffer: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        nonlocal buffer
        content = normalize_text("\n".join(buffer))
        if content:
            sections.append(
                ExtractedSection(title, content, f"lines {start_line}-{end_line}")
            )
        buffer = []

    for line_number, line in enumerate(lines, 1):
        stripped = line.strip()
        if stripped and CHAPTER_PATTERN.match(stripped) and buffer:
            flush(line_number - 1)
            title = stripped
            start_line = line_number + 1
        elif stripped and CHAPTER_PATTERN.match(stripped) and not buffer:
            title = stripped
            start_line = line_number + 1
        elif PAGE_NUMBER_PATTERN.match(stripped):
            continue
        else:
            buffer.append(line)
    flush(len(lines))
    return sections or [ExtractedSection("Text", normalize_text(text), "complete file")]


def _epub_blocks(soup: BeautifulSoup) -> list[str]:
    for element in soup(["script", "style", "nav", "noscript"]):
        element.decompose()
    blocks: list[str] = []
    for element in soup.find_all(["p", "blockquote", "li", "hr"]):
        if element.name == "hr":
            blocks.append("* * *")
            continue
        # Avoid duplicating nested list/blockquote text.
        if element.find_parent(["p", "blockquote", "li"]):
            continue
        value = " ".join(element.get_text(" ", strip=True).split())
        if value:
            blocks.append(value)
    if not blocks:
        value = soup.get_text("\n", strip=True)
        blocks = [line.strip() for line in value.splitlines() if line.strip()]
    return blocks


def _read_epub(path: Path) -> tuple[str, list[ExtractedSection], list[dict[str, Any]]]:
    try:
        book = epub.read_epub(str(path), options={"ignore_ncx": True})
    except Exception as error:  # EbookLib exposes several parser exception types.
        raise IngestError(f"EPUB parsing failed for {path}: {error}") from error
    titles = book.get_metadata("DC", "title")
    book_title = str(titles[0][0]) if titles else path.stem
    document_items = {item.get_id(): item for item in book.get_items_of_type(ITEM_DOCUMENT)}
    ordered_ids = [entry[0] for entry in book.spine]
    items = [document_items[item_id] for item_id in ordered_ids if item_id in document_items]
    if not items:
        items = list(document_items.values())

    sections: list[ExtractedSection] = []
    for position, item in enumerate(items):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        heading = soup.find(["h1", "h2", "h3", "title"])
        chapter_title = (
            " ".join(heading.get_text(" ", strip=True).split())
            if heading is not None
            else Path(item.get_name()).stem
        )
        content = normalize_text("\n\n".join(_epub_blocks(soup)))
        if content:
            # An EPUB spine item can contain several chapters, not just one.
            headings = soup.find_all(["h1", "h2", "h3"])
            if len(headings) <= 1:
                sections.append(ExtractedSection(chapter_title, content, f"spine item {position}: {item.get_name()}"))
            else:
                buffer: list[str] = []
                current_title = chapter_title
                part = 0
                def flush_epub() -> None:
                    nonlocal buffer, part
                    value = normalize_text("\n\n".join(buffer))
                    if value:
                        sections.append(ExtractedSection(current_title, value, f"spine item {position}: {item.get_name()}, heading {part}"))
                        part += 1
                    buffer = []
                for element in soup.find_all(["h1", "h2", "h3", "p", "blockquote", "li", "hr"]):
                    if element.find_parent(["p", "blockquote", "li"]):
                        continue
                    value = " ".join(element.get_text(" ", strip=True).split())
                    if element.name in ("h1", "h2", "h3"):
                        flush_epub()
                        current_title = value or current_title
                    elif element.name == "hr":
                        buffer.append("* * *")
                    elif value:
                        buffer.append(value)
                flush_epub()
    if not sections:
        raise IngestError(f"No readable document content found in EPUB: {path}")
    return book_title, sections, [{"operation": "epub_spine_order", "detail": {"items": len(items)}}]


def _read_docx(path: Path) -> tuple[str, list[ExtractedSection], list[dict[str, Any]]]:
    try:
        document = Document(str(path))
    except Exception as error:
        raise IngestError(f"DOCX parsing failed for {path}: {error}") from error
    core_title = str(document.core_properties.title or "").strip()
    title_source = "core_properties"
    if core_title:
        title = core_title
    else:
        cover_candidates: list[tuple[float, str]] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            runs = [run for run in paragraph.runs if run.text.strip() and run.font.size is not None]
            if value and count_words(value) <= 12 and runs:
                cover_candidates.append((max(run.font.size.pt for run in runs), value))
        largest_size, inferred_title = max(cover_candidates, default=(0.0, ""))
        if largest_size >= 18:
            title = inferred_title
            title_source = f"largest short cover text ({largest_size:g} pt)"
        else:
            title = path.stem
            title_source = "filename"
    sections: list[ExtractedSection] = []
    chapter_title = "Text"
    buffer: list[str] = []
    start_paragraph = 1
    visual_heading_count = 0

    def is_heading(paragraph: Any, value: str) -> tuple[bool, bool]:
        style = paragraph.style.name.lower() if paragraph.style is not None else ""
        if value and (style.startswith("heading") or style in {"title", "subtitle"} or CHAPTER_PATTERN.match(value)):
            return True, False
        # Print-oriented manuscripts often flatten all styles to Normal. A short
        # paragraph whose non-empty runs are uniformly bold and at least 14 pt is
        # still a strong, document-native heading signal.
        runs = [run for run in paragraph.runs if run.text.strip()]
        sizes = [run.font.size.pt for run in runs if run.font.size is not None]
        visual = bool(
            value
            and count_words(value) <= 15
            and runs
            and len(sizes) == len(runs)
            and min(sizes) >= 14
            and all(run.bold is True for run in runs)
        )
        return visual, visual

    def flush(end_paragraph: int) -> None:
        nonlocal buffer
        content = normalize_text("\n\n".join(buffer))
        if content:
            sections.append(
                ExtractedSection(
                    chapter_title,
                    content,
                    f"paragraphs {start_paragraph}-{end_paragraph}",
                )
            )
        buffer = []

    for index, paragraph in enumerate(document.paragraphs, 1):
        value = paragraph.text.strip()
        heading, visual_heading = is_heading(paragraph, value)
        if heading:
            flush(index - 1)
            chapter_title = value
            start_paragraph = index + 1
            visual_heading_count += int(visual_heading)
        elif value:
            buffer.append(value)
    flush(len(document.paragraphs))
    if not sections:
        raise IngestError(f"No body paragraphs found in DOCX: {path}")
    return title, sections, [
        {
            "operation": "docx_paragraph_order",
            "detail": {
                "paragraphs": len(document.paragraphs),
                "visual_headings_detected": visual_heading_count,
                "visual_heading_rule": "all non-empty runs bold, >=14 pt, <=15 words",
                "title_source": title_source,
            },
        }
    ]


def _normal_margin_line(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _read_pdf(
    path: Path,
    *,
    min_words_per_page: int,
    repeated_margin_fraction: float,
) -> tuple[str, list[ExtractedSection], list[dict[str, Any]]]:
    try:
        reader = PdfReader(str(path))
    except Exception as error:
        raise IngestError(f"PDF parsing failed for {path}: {error}") from error
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise IngestError(f"Encrypted PDF requires a password: {path}")
        except Exception as error:
            raise IngestError(f"Encrypted PDF cannot be read: {path}") from error

    page_lines: list[list[str]] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        page_lines.append([line.strip() for line in text.splitlines() if line.strip()])
    word_counts = [count_words("\n".join(lines)) for lines in page_lines]
    low_text_pages = sum(count < min_words_per_page for count in word_counts)
    if not page_lines or low_text_pages / len(page_lines) >= 0.8:
        raise IngestError(
            "PDF appears scanned or has insufficient extractable text. Perform OCR and "
            f"visually verify a sample before import: {path}"
        )

    first_counts = Counter(
        _normal_margin_line(lines[0]) for lines in page_lines if lines
    )
    last_counts = Counter(
        _normal_margin_line(lines[-1]) for lines in page_lines if lines
    )
    threshold = max(3, int(len(page_lines) * repeated_margin_fraction + 0.999))
    repeated_headers = {value for value, count in first_counts.items() if count >= threshold}
    repeated_footers = {value for value, count in last_counts.items() if count >= threshold}

    cleaned_pages: list[str] = []
    removed_margins = 0
    removed_page_numbers = 0
    for lines in page_lines:
        cleaned: list[str] = []
        for position, line in enumerate(lines):
            normalized = _normal_margin_line(line)
            if (position == 0 and normalized in repeated_headers) or (
                position == len(lines) - 1 and normalized in repeated_footers
            ):
                removed_margins += 1
                continue
            if PAGE_NUMBER_PATTERN.match(line):
                removed_page_numbers += 1
                continue
            cleaned.append(line)
        cleaned_pages.append("\n".join(cleaned))
    sections = _split_plain_text("\n\n".join(cleaned_pages))
    title = str(reader.metadata.title) if reader.metadata and reader.metadata.title else path.stem
    log = [
        {
            "operation": "pdf_text_extraction",
            "detail": {
                "pages": len(page_lines),
                "removed_repeated_margins": removed_margins,
                "removed_page_numbers": removed_page_numbers,
                "words_per_page": word_counts,
            },
        }
    ]
    return title, sections, log


def classify_matter(
    title: str,
    index: int,
    total: int,
    *,
    book_title: str | None = None,
    content: str = "",
) -> str:
    normalized = re.sub(r"[^\w\s]", "", title.casefold()).strip()
    front_window = total > 1 and index < max(2, total // 10)
    publishing_markers = (
        "bibliografische information",
        "copyright",
        "impressum",
        "isbn",
        "auflage",
        "books on demand",
        "herstellung und verlag",
    )
    normalized_content = content[:4000].casefold()
    title_matches_book = bool(book_title and _slug(title) == _slug(book_title))
    if normalized in FRONTMATTER_HEADINGS or (
        front_window
        and (
            normalized in {"text", "document"}
            or title_matches_book
            or any(marker in normalized_content for marker in publishing_markers)
        )
    ):
        return "front"
    if normalized in BACKMATTER_HEADINGS or (
        index >= max(0, total - max(2, total // 10)) and normalized in BACKMATTER_HEADINGS
    ):
        return "back"
    return "body"


def extract_book(
    path: str | Path,
    *,
    language: str = "und",
    exclude_frontmatter: bool = True,
    exclude_backmatter: bool = True,
    pdf_min_words_per_page: int = 20,
    repeated_margin_fraction: float = 0.5,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    source = Path(path).resolve()
    extension = source.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise IngestError(f"Unsupported input type {extension!r}: {source}")
    if extension == ".txt":
        title, sections, transformations = _read_txt(source)
    elif extension == ".epub":
        title, sections, transformations = _read_epub(source)
    elif extension == ".docx":
        title, sections, transformations = _read_docx(source)
    else:
        title, sections, transformations = _read_pdf(
            source,
            min_words_per_page=pdf_min_words_per_page,
            repeated_margin_fraction=repeated_margin_fraction,
        )

    source_hash = sha256_file(source)
    book_id = f"{_slug(title)}-{source_hash[:12]}"
    rows: list[dict[str, Any]] = []
    for chapter_index, section in enumerate(sections):
        content = normalize_text(section.content)
        matter = classify_matter(
            section.chapter_title,
            chapter_index,
            len(sections),
            book_title=title,
            content=content,
        )
        included = not (
            (matter == "front" and exclude_frontmatter)
            or (matter == "back" and exclude_backmatter)
        )
        rows.append(
            {
                "book_id": book_id,
                "book_title": title,
                "source_file": str(source),
                "language": language,
                "chapter_index": chapter_index,
                "chapter_title": section.chapter_title,
                "source_position": section.source_position,
                "matter": matter,
                "included": included,
                "word_count": count_words(content),
                "text_sha256": sha256_text(content),
                "content": content,
            }
        )
    if not any(row["included"] and row["word_count"] for row in rows):
        raise IngestError(f"No included body text remains after import: {source}")
    inventory = {
        "book_id": book_id,
        "book_title": title,
        "source_file": str(source),
        "source_extension": extension,
        "source_size_bytes": source.stat().st_size,
        "source_sha256": source_hash,
        "language": language,
        "section_count": len(rows),
        "included_word_count": sum(row["word_count"] for row in rows if row["included"]),
        "imported_at_utc": utc_now(),
    }
    transformation_rows = [
        {
            "book_id": book_id,
            "source_file": str(source),
            "timestamp_utc": utc_now(),
            "operation": "unicode_normalization",
            "detail": {"form": "NFC"},
        },
        *(
            {
                "book_id": book_id,
                "source_file": str(source),
                "timestamp_utc": utc_now(),
                **entry,
            }
            for entry in transformations
        ),
        {
            "book_id": book_id,
            "source_file": str(source),
            "timestamp_utc": utc_now(),
            "operation": "matter_classification",
            "detail": {
                "frontmatter_excluded": exclude_frontmatter,
                "backmatter_excluded": exclude_backmatter,
                "section_matter": [row["matter"] for row in rows],
            },
        },
    ]
    return inventory, rows, transformation_rows


def discover_books(input_dir: str | Path) -> list[Path]:
    directory = project_path(input_dir)
    if not directory.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {directory}")
    return sorted(
        path for path in directory.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def ingest_books(
    input_dir: str | Path,
    output_dir: str | Path,
    *,
    language: str = "und",
    exclude_frontmatter: bool = True,
    exclude_backmatter: bool = True,
    pdf_min_words_per_page: int = 20,
    repeated_margin_fraction: float = 0.5,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    sources = discover_books(input_dir)
    if not sources:
        raise FileNotFoundError(
            f"No supported EPUB, DOCX, TXT, or PDF files found in {project_path(input_dir)}"
        )
    inventories: list[dict[str, Any]] = []
    sections: list[dict[str, Any]] = []
    transformations: list[dict[str, Any]] = []
    for source in sources:
        inventory, book_sections, book_transformations = extract_book(
            source,
            language=language,
            exclude_frontmatter=exclude_frontmatter,
            exclude_backmatter=exclude_backmatter,
            pdf_min_words_per_page=pdf_min_words_per_page,
            repeated_margin_fraction=repeated_margin_fraction,
        )
        inventories.append(inventory)
        sections.extend(book_sections)
        transformations.extend(book_transformations)

    inventory_frame = pd.DataFrame(inventories).sort_values("book_id").reset_index(drop=True)
    section_frame = pd.DataFrame(sections).sort_values(["book_id", "chapter_index"]).reset_index(drop=True)
    output = project_path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    inventory_frame.to_parquet(output / "book_inventory.parquet", index=False)
    section_frame.to_parquet(output / "book_sections.parquet", index=False)
    with (output / "transformations.jsonl").open("w", encoding="utf-8") as handle:
        for record in transformations:
            handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")
    return inventory_frame, section_frame
