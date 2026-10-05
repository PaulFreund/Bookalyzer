"""Offline checks of frozen assets and dynamic imports. Uses synthetic text only."""
import base64
import json
import re
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd

from .desktop import DesktopService, DEFAULTS
from .provenance import assert_storyscope_immutable


def main():
    from .codex_adapter import _load_stage5
    from .config import PROJECT_ROOT
    assert_storyscope_immutable()
    stage5 = _load_stage5(PROJECT_ROOT / "vendor/storyscope")
    taxonomy = stage5.Taxonomy.from_json(str(PROJECT_ROOT / "data/reference/storyscope/taxonomy.json"))
    assert len(taxonomy.dimensions) == 10
    with TemporaryDirectory(prefix="Bookalyzer-offline-") as folder:
        root = Path(folder)
        service = DesktopService(root / "library", seed=False)
        try:
            _, specs = service._taxonomy()
            documents = []
            for name, value in (("Referenz", 1), ("Vergleich", 5)):
                text = ("Kapitel 1\n" + (name + " im ersten Text. ") * 35 +
                        "\nKapitel 2\n" + (name + " im zweiten Text. ") * 35)
                doc = service.prepare(name + ".txt", base64.b64encode(text.encode()).decode(),
                                      {**DEFAULTS, "detectChapters": True})
                directory = service._directory(doc["id"])
                segments = pd.read_parquet(directory / "segments.parquet")
                rows = []
                for index, segment in segments.iterrows():
                    row = {key: (re.match(r"^\d+", spec.values[0]).group() if spec.feature_type == "scale" else spec.values[0])
                           for key, spec in specs.items()}
                    row.update(segment.drop(labels=["human_story"]).to_dict())
                    row.update(TMP_ORD_010=str(value + (index if value == 1 else 0)), extraction_model="offline-fixture")
                    rows.append(row)
                features = pd.DataFrame(rows)
                features.to_parquet(directory / "features.parquet", index=False)
                service._complete(doc, segments, features)
                documents.append(doc)
            baseline = service.baseline_create(documents[0]["id"])
            result = service.baseline_compare([documents[1]["id"]], baseline["id"])
            assert result["rows"][0]["distance"] > 0
            assert service.compare([d["id"] for d in documents])["available"]
            quality = service.quality([documents[1]["id"]], baselineId=baseline["id"])
            assert len(quality["catalog"]["metrics"]) == 51
            assert quality["rows"][0]["quality"]["values"]["mattr"] is not None
            assert len(service.quality([], bookId=documents[1]["id"])["rows"]) == 2
            assert all(c["narrativeComplete"] for c in service.chapters(documents[1]["id"], baseline["id"])["chapters"])
            # Exercise import resources that are commonly missed by freezers.
            from docx import Document
            from ebooklib import epub
            from pypdf import PdfWriter
            from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
            body = "The wind moved through the quiet garden. " * 30
            word = Document()
            word.add_heading("Chapter 1", level=1)
            word.add_paragraph(body)
            word.save(root / "sample.docx")
            book = epub.EpubBook()
            book.set_identifier("runtime-fixture")
            book.set_title("Runtime")
            book.set_language("en")
            chapter = epub.EpubHtml(title="Chapter 1", file_name="text.xhtml", lang="en")
            chapter.content = "<h1>Chapter 1</h1><p>" + body + "</p>"
            for item in (chapter, epub.EpubNcx(), epub.EpubNav()):
                book.add_item(item)
            book.spine = [chapter]
            book.toc = (chapter,)
            epub.write_epub(str(root / "sample.epub"), book)
            pdf = PdfWriter()
            page = pdf.add_blank_page(600, 800)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(("BT /F1 12 Tf 40 740 Td (" + body + ") Tj ET").encode("ascii"))
            page[NameObject("/Contents")] = stream
            pdf.write(root / "sample.pdf")
            for ext in ("docx", "epub", "pdf"):
                prepared = service.prepare("sample." + ext, base64.b64encode((root / ("sample." + ext)).read_bytes()).decode(), DEFAULTS)
                assert prepared["wordCount"] > 100
            print(json.dumps({"ok": True, "formats": ["txt", "docx", "epub", "pdf"],
                              "features": len(specs), "baselineDistance": result["rows"][0]["distance"]}), flush=True)
        finally:
            service.executor.shutdown()
