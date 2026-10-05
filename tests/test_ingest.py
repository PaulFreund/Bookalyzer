from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from ebooklib import epub
from reportlab.pdfgen import canvas
from docx.shared import Pt

from bookanalyzer.ingest import IngestError, extract_book


BODY = " ".join(f"word{i}" for i in range(45))


def test_epub_detects_multiple_chapters_inside_one_spine_item(tmp_path):
    path = tmp_path / "combined.epub"
    book = epub.EpubBook()
    book.set_identifier("combined-fixture")
    book.set_title("Novel")
    book.set_language("de")
    page = epub.EpubHtml(title="Novel", file_name="all.xhtml", lang="de")
    page.content = "<h1>Kapitel 1</h1><p>First text remains.</p><h1>Kapitel 2</h1><p>Second text remains.</p>"
    book.add_item(page)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (page,)
    book.spine = [page]
    epub.write_epub(str(path), book)
    _, sections, _ = extract_book(path)
    assert [s["chapter_title"] for s in sections] == ["Kapitel 1", "Kapitel 2"]
    assert [s["content"] for s in sections] == ["First text remains.", "Second text remains."]


def test_txt_import_normalizes_unicode_and_keeps_body(tmp_path: Path) -> None:
    path = tmp_path / "novel.txt"
    path.write_text(f"CHAPTER 1\n\nCafe\u0301 begins.\n\n{BODY}", encoding="utf-8")
    inventory, sections, transformations = extract_book(path, language="en")
    assert inventory["language"] == "en"
    assert sections[0]["content"].startswith("Café")
    assert sections[0]["text_sha256"]
    assert any(item["operation"] == "unicode_normalization" for item in transformations)


def test_docx_import_uses_heading_boundaries(tmp_path: Path) -> None:
    path = tmp_path / "novel.docx"
    document = Document()
    document.core_properties.title = "DOCX Novel"
    document.add_heading("Chapter One", level=1)
    document.add_paragraph(BODY)
    document.add_heading("Chapter Two", level=1)
    document.add_paragraph(BODY)
    document.save(path)
    inventory, sections, _ = extract_book(path)
    assert inventory["book_title"] == "DOCX Novel"
    assert [row["chapter_title"] for row in sections] == ["Chapter One", "Chapter Two"]


def test_docx_import_recovers_visual_headings_and_frontmatter(tmp_path: Path) -> None:
    path = tmp_path / "visual.docx"
    document = Document()
    cover = document.add_paragraph()
    cover_run = cover.add_run("Visual Novel")
    cover_run.bold = True
    cover_run.font.size = Pt(30)
    document.add_paragraph("Copyright and ISBN 123")
    for heading_text in ("The First Door", "The Second Door"):
        heading = document.add_paragraph()
        run = heading.add_run(heading_text)
        run.bold = True
        run.font.size = Pt(16)
        document.add_paragraph(BODY)
    document.save(path)
    inventory, sections, _ = extract_book(path)
    assert inventory["book_title"] == "Visual Novel"
    assert [row["chapter_title"] for row in sections] == [
        "Visual Novel",
        "The First Door",
        "The Second Door",
    ]
    assert sections[0]["matter"] == "front"
    assert sections[0]["included"] is False
    assert all(row["included"] for row in sections[1:])


def test_epub_import_follows_spine(tmp_path: Path) -> None:
    path = tmp_path / "novel.epub"
    book = epub.EpubBook()
    book.set_identifier("fixture")
    book.set_title("EPUB Novel")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter One", file_name="chapter.xhtml", lang="en")
    chapter.content = f"<html><body><h1>Chapter One</h1><p>{BODY}</p></body></html>"
    book.add_item(chapter)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    epub.write_epub(path, book)
    inventory, sections, _ = extract_book(path)
    assert inventory["book_title"] == "EPUB Novel"
    assert sections[0]["chapter_title"] == "Chapter One"
    assert "word44" in sections[0]["content"]


def test_text_pdf_import_and_scanned_pdf_refusal(tmp_path: Path) -> None:
    text_path = tmp_path / "text.pdf"
    document = canvas.Canvas(str(text_path))
    text = document.beginText(50, 780)
    for line_start in range(0, 45, 9):
        text.textLine(" ".join(f"word{i}" for i in range(line_start, line_start + 9)))
    document.drawText(text)
    document.save()
    inventory, sections, _ = extract_book(text_path, pdf_min_words_per_page=20)
    assert inventory["included_word_count"] >= 40
    assert sections

    scanned_path = tmp_path / "scanned.pdf"
    blank = canvas.Canvas(str(scanned_path))
    blank.showPage()
    blank.save()
    with pytest.raises(IngestError, match="OCR"):
        extract_book(scanned_path, pdf_min_words_per_page=20)
